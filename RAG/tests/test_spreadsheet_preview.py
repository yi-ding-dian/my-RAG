"""类 Excel 预览 HTML 渲染测试（spreadsheet_preview）

覆盖：HTML 骨架（纯 CSS tab 切换无 script）/ 多 sheet /
合并单元格 rowspan-colspan / 列宽 / 百分比与日期格式 /
值转义（防 XSS）/ csv 与 xls 入口 / 空格表报错。
"""
from __future__ import annotations

import re

import pytest


def _make_xlsx(tmp_path, name="p.xlsx"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据表"
    ws.append(["城市", "预算", "占比", "日期"])
    ws.append(["桂林", 3000, 0.125, "2026-08-22"])
    ws.merge_cells("A2:A3")
    ws.cell(row=2, column=3).number_format = "0.0%"  # 百分比显示
    ws.append(["长沙", 1500, 0.5, "2026-08-23"])
    ws3 = wb.create_sheet("备注")
    ws3.append(["说明", "内容"])
    ws3.append(["安全", "<script>alert(1)</script>"])
    p = tmp_path / name
    wb.save(p)
    return p


def test_xlsx_html_structure_and_tabs(tmp_path):
    from backend.services.spreadsheet_preview import render_spreadsheet_html
    html = render_spreadsheet_html(_make_xlsx(tmp_path))
    assert html.startswith("<!doctype html>")
    assert "<script" not in html, "纯 CSS 切换，零 JS"
    assert "input#sheet-0:checked ~ .sheet.sheet-0" in html
    assert 'label for="sheet-0"' in html and 'label for="sheet-1"' in html
    assert html.count("<table>") == 2


def test_xlsx_merged_and_width_and_format(tmp_path):
    from backend.services.spreadsheet_preview import render_spreadsheet_html
    html = render_spreadsheet_html(_make_xlsx(tmp_path))
    assert 'rowspan="2"' in html, "合并单元格 rowspan 还原"
    assert "width:" in html and "px" in html, "列宽还原"
    assert "12.5%" in html, "百分比格式还原"
    assert "2026-08-22" in html, "日期文本保留"


def test_xlsx_value_escaped(tmp_path):
    from backend.services.spreadsheet_preview import render_spreadsheet_html
    html = render_spreadsheet_html(_make_xlsx(tmp_path))
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html, "单元格值须转义（防 XSS）"


def test_csv_html(tmp_path):
    from backend.services.spreadsheet_preview import render_spreadsheet_html
    p = tmp_path / "t.csv"
    p.write_bytes("城市,预算\n桂林,3000\n".encode("gbk"))
    html = render_spreadsheet_html(p)
    assert "城市" in html and "桂林" in html
    assert html.count("<table>") == 1


def test_empty_sheet_raises(tmp_path):
    from backend.services.spreadsheet_preview import render_spreadsheet_html
    import openpyxl
    p = tmp_path / "e.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append([None, None])
    wb.save(p)
    with pytest.raises(RuntimeError, match="无可用工作表"):
        render_spreadsheet_html(p)


def test_xls_basic(tmp_path):
    import xlwt
    from backend.services.spreadsheet_preview import render_spreadsheet_html
    wb = xlwt.Workbook()
    sh = wb.add_sheet("Old")
    sh.write(0, 0, "项目")
    sh.write(0, 1, "值")
    sh.write(1, 0, "A")
    sh.write(1, 1, 123.0)
    sh.write_merge(2, 2, 0, 1, "合并行")  # row2 合并两列
    p = tmp_path / "old.xls"
    wb.save(str(p))
    html = render_spreadsheet_html(p)
    assert 'label for="sheet-0"' in html
    assert "合并行" in html
    assert 'colspan="2"' in html
