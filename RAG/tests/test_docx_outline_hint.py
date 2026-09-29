"""docx-outline 层级探测测试：前端「选了 MinerU 会丢层级」黄条的判据

背景（实测事故）：一份有规范标题样式的 .docx（Word 里 496 个 Heading 1~6）
被 MinerU 解析后，544 个标题里 541 个被压成同一级、还多出 49 个把正文
（"注意事项:"）误标成标题的假标题，切块详情的「目录」因此不可读；同一份
文档换「结构解析」（docx_struct）则层级完整。

前端 ParseConfigModal 在「解析方式=MinerU + Word 文档」时会调本接口探测，
判据 **max_level >= 2 且 count >= 3** 就亮黄条。本文件验证这个判据在真实
docx 上成立（有层级→亮；无层级→不亮，避免误报）。
"""
from __future__ import annotations

import io

from conftest import create_kb, upload_doc

DOCX_MIME = ("application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document")


def _docx_bytes(headings: bool) -> bytes:
    """造一个 docx：headings=True 用 Word 标题样式，False 全是普通正文"""
    from docx import Document

    d = Document()
    if headings:
        d.add_heading("第一章 概述", level=1)
        d.add_paragraph("正文内容。")
        d.add_heading("1.1 背景", level=2)
        d.add_paragraph("正文内容。")
        d.add_heading("1.1.1 细节", level=3)
        d.add_paragraph("正文内容。")
        d.add_heading("第二章 安装", level=1)
        d.add_paragraph("正文内容。")
    else:
        for i in range(8):
            d.add_paragraph(f"这是第 {i} 段普通正文，没有任何标题样式。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _outline(client, kb_id, doc_id, headers):
    return client.get(
        f"/api/kbs/{kb_id}/documents/{doc_id}/docx-outline", headers=headers)


def _upload(client, kb_id, name, data, headers):
    return upload_doc(client, kb_id, filename=name, content=data,
                      mime=DOCX_MIME, headers=headers)


class TestOutlineProbe:

    def test_multilevel_docx_reports_depth(self, client, admin_headers):
        """★ 有 Heading 1~3 的规范 docx → count>=3 且 max_level>=2

        这两条同时成立，前端才会亮「MinerU 会压平层级」的黄条。
        """
        kb = create_kb(client)
        doc = _upload(client, kb["id"], "规范.docx", _docx_bytes(True),
                      admin_headers)
        resp = _outline(client, kb["id"], doc["id"], admin_headers)
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["warning"] is None
        assert data["count"] >= 3, f"应读出多个标题，实际 {data['count']}"
        assert data["max_level"] >= 2, \
            f"应读出多级层级，实际 max_level={data['max_level']}"
        # 前端判据原样断言（条件变了这里会红）
        assert data["max_level"] >= 2 and data["count"] >= 3, \
            "前端据此亮黄条：有层级的 docx 必须命中判据"

    def test_plain_docx_does_not_trigger(self, client, admin_headers):
        """全是普通正文的 docx → max_level < 2（不该亮黄条，避免误报）"""
        kb = create_kb(client)
        doc = _upload(client, kb["id"], "纯正文.docx", _docx_bytes(False),
                      admin_headers)
        resp = _outline(client, kb["id"], doc["id"], admin_headers)
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert not (data["max_level"] >= 2 and data["count"] >= 3), \
            "无标题层级的文档不该命中判据（否则提示会变成噪音）"

    def test_non_word_document_400(self, client, admin_headers):
        """非 Word 文档 → 400（前端只在 docx/doc 时调）"""
        kb = create_kb(client)
        doc = upload_doc(client, kb["id"], filename="普通.txt")
        resp = _outline(client, kb["id"], doc["id"], admin_headers)
        assert resp.status_code == 400

    def test_unknown_doc_404(self, client, admin_headers):
        kb = create_kb(client)
        resp = _outline(client, kb["id"], "nonexist", admin_headers)
        assert resp.status_code == 404
