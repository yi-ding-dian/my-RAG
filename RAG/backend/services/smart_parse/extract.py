"""轻量本地文本提取 + docx 结构探测（本包唯一 IO 层）

**不调 MinerU/DeepDoc**：只用本地能力（直读 / pypdf / python-docx / LibreOffice
转换）拿到"够画像用"的文本与结构。真正的解析在入库阶段由 ingestion 走解析器。

- txt/md 直读（utf-8，GBK 兼容回退）
- pdf 走 parsers.client._extract_plain（pypdf）
- docx/doc 走结构化读取：一次 OOXML 解析同时拿到「纯文本 + 规范性探测」
  （doc 先经 LibreOffice 转 docx，仅转一次，临时目录即用即清）
"""
from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path

from backend.chunking import detect_heading_systems
from backend.services.parsers.client import get_parser_client
from backend.services.smart_parse.types import (_DOC_LIKE_FILE_TYPES,
                                                _PARSER_FILE_TYPES,
                                                _TEXT_FILE_TYPES)

logger = logging.getLogger(__name__)

# 规范性判定阈值（样式标题数量 + 非空段落占比）：达到即认为文档有规范写作
# 格式（标题层级明确），建议本地结构化解析（docx_struct）保留层级
NORMATIVE_MIN_STYLE_HEADINGS = 3
NORMATIVE_MIN_STYLE_RATIO = 0.02

# 编号体系命中统计的标题样本上限（探测开销封顶，样本够判断体系即可）
_HEADING_SAMPLE_LIMIT = 200


def extract_text(path: Path, file_type: str) -> tuple[str, bool, str | None]:
    """轻量本地文本提取：txt/md 直读；pdf/docx 复用 parsers.client._extract_plain。

    返回 (text, extracted, warning)：extracted=False 时 text 为空串。
    """
    try:
        if file_type in _TEXT_FILE_TYPES:
            if not path.exists():
                return "", False, "源文件缺失，无法提取文本"
            raw = path.read_bytes()
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = raw.decode("gbk", errors="replace")
        elif file_type in _PARSER_FILE_TYPES:
            if not path.exists():
                return "", False, "源文件缺失，无法提取文本"
            text = get_parser_client()._extract_plain(path, file_type)
        else:
            return "", False, f"暂不支持该类型（{file_type}）画像分析"
    except Exception as e:
        logger.warning("智能解析文本提取失败 %s: %s", path.name, e)
        return "", False, f"文本提取失败: {e}"
    if not text.strip():
        return "", False, "未提取到文本内容（文件为空或为扫描件）"
    return text.strip(), True, None


def outline_level(paragraph) -> int | None:
    """段落大纲层级：pPr/outlineLvl 显式层级 或 样式名 Heading N/标题 N；
    非标题段落返回 None（访问异常防御式返回 None，不影响整体探测）"""
    try:
        p_pr = paragraph._p.pPr
        if p_pr is not None and p_pr.outlineLvl is not None:
            return int(p_pr.outlineLvl.val)
        name = (paragraph.style.name if paragraph.style is not None
                else "") or ""
    except Exception:
        return None
    m = re.match(r"^(?:heading|标题)\s*([1-9])$", name.strip(), re.I)
    return int(m.group(1)) - 1 if m else None


def has_num_pr(paragraph) -> bool:
    """段落是否挂自动编号/多级列表（pPr.numPr）：规范文档的标题常由多级
    列表编号（如"一、"/"1.1"），没有 Heading 样式也可视为结构化标题"""
    try:
        p_pr = paragraph._p.pPr
        return p_pr is not None and p_pr.numPr is not None
    except Exception:
        return False


def probe_docx_structure(document) -> dict:
    """docx 结构探测（规范性判定；python-docx 轻量读取，不调 MinerU/DeepDoc）

    - style_headings：样式标题段落数（outlineLvl 或 Heading 样式——Word 里
      "这是标题"的显式标记）；
    - numbered_headings：挂自动编号（numPr）的段落数，仅作画像参考、**不参与**
      规范性判据：正文枚举列表（"1）…"）同样带 numPr，计入会把列表密集的
      普通文档误判为规范文档；
    - style_ratio：样式标题 / 非空段落数（样式标题覆盖率）；
    - style_heading_systems：样式标题的编号体系命中（复用 heading_presets
      .detect_heading_systems，与切块层级推断同口径）；
    - is_normative：样式标题 >= 3 个且覆盖率 >= 2% → 有规范写作格式
      （适合结构化解析保留标题层级）。
    """
    total = 0
    style_headings = 0
    numbered_headings = 0
    titles: list = []
    for p in document.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        total += 1
        if outline_level(p) is not None:
            style_headings += 1
            if len(titles) < _HEADING_SAMPLE_LIMIT:
                titles.append(text)
        elif has_num_pr(p):
            numbered_headings += 1
    ratio = style_headings / total if total else 0.0
    return {
        "style_headings": style_headings,
        "numbered_headings": numbered_headings,
        "total_paragraphs": total,
        "style_ratio": round(ratio, 4),
        "style_heading_systems": detect_heading_systems(titles),
        "is_normative": (style_headings >= NORMATIVE_MIN_STYLE_HEADINGS
                         and ratio >= NORMATIVE_MIN_STYLE_RATIO),
    }


def extract_docx_like(path: Path, file_type: str) -> tuple[str, dict, str | None]:
    """docx/doc 结构化读取（画像用）：一次 OOXML 解析出「纯文本 + 结构探测」

    - docx 直接读；doc 先经 LibreOffice 转 docx（仅转一次，临时目录即用即清，
      见 parsers.client.convert_doc_to_docx）；
    - 文本口径与 parsers.client._extract_plain 一致（段落 + 表格行拼接）；
    - 结构探测见 probe_docx_structure（规范性判定，引擎建议用）。
    返回 (text, probe, warning)：warning 非空 = 提取失败（text 为空串）。
    """
    tmp_dir = None
    try:
        if file_type == "doc":
            from backend.services.parsers.client import convert_doc_to_docx
            path, tmp_dir = convert_doc_to_docx(path)
        import docx  # python-docx
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    paras.append(" | ".join(cells))
        text = "\n\n".join(paras).strip()
        probe = probe_docx_structure(d)
    except Exception as e:
        logger.warning("智能解析结构化读取失败 %s: %s", path.name, e)
        return "", {}, f"文本提取失败: {e}"
    finally:
        if tmp_dir is not None:
            shutil.rmtree(tmp_dir, ignore_errors=True)
    if not text:
        return "", probe, "未提取到文本内容（文件为空或为扫描件）"
    return text, probe, None


async def extract_for_profile(path: Path, file_type: str) -> tuple[str, bool, dict, str | None]:
    """按文件类型分派提取（docx/doc 走 to_thread，避免阻塞事件循环）

    返回 (text, extracted, docx_probe, warning)；非 docx/doc 时 docx_probe 为 {}。
    """
    if file_type in _DOC_LIKE_FILE_TYPES:
        import asyncio
        text, probe, warning = await asyncio.to_thread(
            extract_docx_like, path, file_type)
        return text, bool(text.strip()), probe, warning
    text, extracted, warning = extract_text(path, file_type)
    return text, extracted, {}, warning
