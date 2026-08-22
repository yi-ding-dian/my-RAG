"""xlsx 读取器（openpyxl，值直读）

- load_workbook(data_only=True)：公式取缓存值（Excel/WPS 保存均写缓存；
  从未被应用打开的公式单元格返回 None → ""，可接受，见模块 docstring）；
- 合并单元格：openpyxl ranges（1 基，含端点）→ 统一 0 基闭区间交
  spreadsheet_reader.fill_merged（与 xls 共用填充逻辑）；
- 单元格格式：None → ""，datetime/date/time → ISO 文本（可读、可检索），
  数字/字符串原样转文本；布尔 → "TRUE"/"FALSE"；
- 空 sheet（所有行全空）跳过；多 sheet 全部导入（"## Sheet: 名称" 分节）。
"""
from __future__ import annotations

import datetime
from pathlib import Path
from typing import List

from backend.services.spreadsheet_reader import (Sheet, fill_merged,
                                                 pad_rows)


def _cell_text(value) -> str:
    """单元格值 → 文本（日期转 ISO、布尔转文本、None → 空）"""
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime.datetime) \
            else value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_xlsx(path: Path) -> List[Sheet]:
    """读取工作簿全部工作表中非空 sheet（空 sheet/全空行跳过）"""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    sheets: List[Sheet] = []
    for ws in wb.worksheets:
        rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row]
            # 全空行跳过（Excel 尾部空行/格式化残留行）
            if not any(c for c in cells):
                continue
            rows.append(cells)
        if not rows:
            continue
        # 合并单元格：openpyxl 1 基含端点 → 0 基闭区间（min_row-1, min_col-1…）
        ranges = [
            (rng.min_row - 1, rng.min_col - 1,
             rng.max_row - 1, rng.max_col - 1)
            for rng in ws.merged_cells.ranges
        ]
        fill_merged(rows, ranges)
        sheets.append(Sheet(name=ws.title, rows=pad_rows(rows),
                            stats={"merged_cells": len(ranges)}))
    return sheets
