"""Excel/CSV 原始观感预览：读取文件结构与常用样式，渲染"类 Excel"HTML。

设计（对齐 WPS/Excel 打开观感，零新增依赖、离线内网可用）：
- 还原：网格线、**合并单元格**（rowspan/colspan）、列宽/行高、
  字体（加粗/颜色/字号）、背景色、对齐、百分比/日期显示格式；
- **多 sheet 标签切换用纯 CSS radio**（内置脚本零 JS，无脚本注入面）；
- 数值全部 html.escape，样式取值不可靠时不虚构（尽力而为）；
- xlsx 走 openpyxl（data_only=True，公式取缓存值，column_dimensions/
  row_dimensions/merged_cells/cell 样式全支持）；xls 走 xlrd
  （formatting_info=True，列宽/合并/字体/底色尽力提取）；csv 简单网格。
- 不还原：图表/图片/条件格式/批注（内网报表 90% 场景不受影响）。

前端预览弹窗 iframe 展示本模块产物（/raw 对 excel 返回 text/html）。
"""
from __future__ import annotations

import csv
import datetime
import html
import io
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# ---- 通用 HTML 骨架（样式 + 纯 CSS tab 切换） ----

_CSS = """
/* 主题变量（iframe：前端按宿主主题注入 html data-theme 属性） */
:root {
  --sp-page-bg: #fff; --sp-page-fg: #1f2328;
  --sp-tabs-bg: #f6f8fa; --sp-border: #d0d7de;
  --sp-tab-bg: #eef1f6; --sp-tab-fg: #1f2328;
  --sp-tab-active-bg: #ffffff; --sp-tab-active-fg: #111827;
}
[data-theme='dark'] {
  --sp-page-bg: #1b2028; --sp-page-fg: #e5e7eb;
  --sp-tabs-bg: #171c23; --sp-border: #2e3644;
  --sp-tab-bg: #232a35; --sp-tab-fg: #c9d1d9;
  --sp-tab-active-bg: #2a3240; --sp-tab-active-fg: #f3f4f6;
}
* { box-sizing: border-box; margin: 0; padding: 0; }
body { font-family: "Microsoft YaHei", "PingFang SC", system-ui, sans-serif;
       background: var(--sp-page-bg); color: var(--sp-page-fg); }
.tabs { position: sticky; top: 0; z-index: 10; display: flex; flex-wrap: wrap;
        gap: 4px; padding: 8px 10px 0; background: var(--sp-tabs-bg);
        border-bottom: 1px solid var(--sp-border); }
.tabs label { padding: 5px 14px; font-size: 13px; border-radius: 6px 6px 0 0;
              background: var(--sp-tab-bg); color: var(--sp-tab-fg);
              cursor: pointer; user-select: none;
              border: 1px solid var(--sp-border); border-bottom: none; }
/* radio 在 body 直接子级（tabs 之外），按 name 隐藏所有选择圆点 */
input[name="sh"] { display: none; }
.sheet { display: none; padding: 10px; }
input#sheet-0:checked ~ .sheet.sheet-0 { display: block; }
input#sheet-1:checked ~ .sheet.sheet-1 { display: block; }
input#sheet-2:checked ~ .sheet.sheet-2 { display: block; }
input#sheet-3:checked ~ .sheet.sheet-3 { display: block; }
input#sheet-4:checked ~ .sheet.sheet-4 { display: block; }
input#sheet-5:checked ~ .sheet.sheet-5 { display: block; }
input#sheet-6:checked ~ .sheet.sheet-6 { display: block; }
input#sheet-7:checked ~ .sheet.sheet-7 { display: block; }
input#sheet-8:checked ~ .sheet.sheet-8 { display: block; }
input#sheet-9:checked ~ .sheet.sheet-9 { display: block; }
/* 选中的 tab：radio 与 label 不同级，用 ~ + nth-child 定位（10 个上限） */
input#sheet-0:checked ~ .tabs label:nth-child(1),
input#sheet-1:checked ~ .tabs label:nth-child(2),
input#sheet-2:checked ~ .tabs label:nth-child(3),
input#sheet-3:checked ~ .tabs label:nth-child(4),
input#sheet-4:checked ~ .tabs label:nth-child(5),
input#sheet-5:checked ~ .tabs label:nth-child(6),
input#sheet-6:checked ~ .tabs label:nth-child(7),
input#sheet-7:checked ~ .tabs label:nth-child(8),
input#sheet-8:checked ~ .tabs label:nth-child(9),
input#sheet-9:checked ~ .tabs label:nth-child(10) {
  background: var(--sp-tab-active-bg); color: var(--sp-tab-active-fg);
  font-weight: 600;
}
table { border-collapse: collapse; table-layout: auto; }
td, th { border: 1px solid var(--sp-border); padding: 2px 6px;
         font-size: 13px; line-height: 1.5; overflow: hidden;
         text-overflow: ellipsis; white-space: nowrap; }
.table-scroll { overflow: auto; max-width: 100%; }
"""

_HEAD_TEMPLATE = """<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>{css}</style></head><body>
{tabs}
{sheets}
</body></html>
"""


def _sheet_tabs_html(names: List[str]) -> str:
    """radio 必须为 body 直接子级（与 .sheet 同级，CSS 兄弟选择器才生效）；
    label 打包进 .tabs 栏（label for 关联 radio，位置不限）"""
    inputs: List[str] = []
    labels: List[str] = []
    for i, name in enumerate(names[:10]):
        checked = " checked" if i == 0 else ""
        inputs.append(f'<input type="radio" id="sheet-{i}" name="sh"{checked}>')
        labels.append(f'<label for="sheet-{i}">{html.escape(name)}</label>')
    return ("\n".join(inputs) + '\n<div class="tabs">\n'
            + "\n".join(labels) + "\n</div>")


def _wrap_sheet(name: str, table_html: str, idx: int) -> str:
    return (f'<div class="sheet sheet-{idx}">\n'
            f'<div class="table-scroll">\n{table_html}\n</div>\n</div>')


def _apply_number_format(value: float, fmt: str) -> str:
    """Excel number_format → 展示文本（货币/千分位/固定小数/百分比，尽力还原）

    覆盖常见企业内部报表格式（如 ¥#,##0.00、#,##0.0、0.00%、$#,##0）；
    未识别或异常 → 原样值回退，不虚构格式。注意: 数字格式里带引号的
    文本片段(如 "¥"#,##0.00)按同规则处理，先取货币符号再应用数值格式。
    """
    fmt = fmt or ""
    try:
        # GENERAL / 无格式 / 文本格式：原样输出（整数去尾 .0；小数净化尾差）
        if not fmt or fmt.upper() in ("GENERAL", "General") or "@" in fmt:
            if value == int(value):
                return str(int(value))
            from backend.services.spreadsheet.reader import humanize_number
            return humanize_number(value)
        # 货币符号(带引号或裸符号)
        symbol = ""
        m = re.search(r'"?([¥￥$€£])"?', fmt)
        if m:
            symbol = m.group(1)
        dec = 2
        if ".00" in fmt:
            dec = 2
        elif ".0" in fmt:
            dec = 1
        elif "." in fmt and not symbol:
            dec = 2
        else:
            dec = 0
        # 有 / 无千分位
        has_comma = "," in fmt
        if has_comma:
            num = f"{value:,.{dec}f}"
        else:
            num = f"{value:.{dec}f}"
        if "%" in fmt:
            num = f"{value * 100:.{max(dec, 0)}f}%"
            symbol = ""
        return f"{symbol}{num}"
    except Exception:
        return str(value)


def _fmt_value(value, number_format: str = "") -> str:
    """单元格值展示文本（数字格式感知：货币/千分位/百分比/日期）"""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        if number_format.lower().find("yyyy") != -1 or "m/d" in number_format.lower():
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return _apply_number_format(float(value), number_format)
    return str(value)


# ---- xlsx（openpyxl） ----

def _xlsx_color(font_or_fill) -> Optional[str]:
    """openpyxl 颜色对象 → css 颜色（theme 色无能力还原，仅实值色）"""
    try:
        c = getattr(font_or_fill, "rgb", None)
        if isinstance(c, str) and c.startswith("#"):
            return c
        if isinstance(c, str) and len(c) == 8 and c.lower().startswith("ff"):
            return f"#{c[2:]}"
        return None
    except Exception:
        return None


def _render_xlsx(path: Path) -> Tuple[List[str], List[str]]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        wb_formula = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        wb_formula = None
    # 公式计算引擎（有公式时重算；缓存命中秒读；失败回退公式文本）
    _recalc: dict = {}
    if wb_formula is not None:
        try:
            has_formula = any(
                isinstance(c.value, str) and c.value.startswith("=")
                for sheet in wb_formula.worksheets for row in sheet.iter_rows()
                for c in row
            )
            if has_formula:
                from backend.services.spreadsheet.formula import recalc_cells
                _recalc = recalc_cells(path)
        except Exception:
            _recalc = {}
    names: List[str] = []
    sheet_htmls: List[str] = []
    for ws in wb.worksheets[:10]:
        names.append(ws.title)
        # 列宽/行高（字符宽 → 像素近似：7px/字符 + padding）
        col_w: Dict[int, float] = {}
        for key, dim in ws.column_dimensions.items():
            if dim and dim.width:
                col_w[key] = max(40.0, round(dim.width * 7) + 12)
        row_h: Dict[int, float] = {}
        for r, dim in ws.row_dimensions.items():
            if dim and dim.height:
                row_h[r] = round(float(dim.height)) + 4
        # 合并单元格（左上坐标 → 跨行跨列；其余坐标跳过输出）
        merge_span: Dict[Tuple[int, int], Tuple[int, int]] = {}
        merged_covered = set()
        for rng in ws.merged_cells.ranges:
            merge_span[(rng.min_row, rng.min_col)] = (
                rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        merged_covered.add((r, c))
        # 空 sheet（无任何值）跳过
        if not any(cell.value not in (None, "") for row in ws.iter_rows()
                   for cell in row):
            names.pop()
            continue
        rows_html: List[str] = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                min_col=1, max_col=ws.max_column):
            cells: List[str] = []
            for cell in row:
                coord = (cell.row, cell.column)
                if coord in merged_covered:
                    continue
                attrs: List[str] = []
                col_letter = get_column_letter(cell.column)
                w = int(col_w.get(col_letter, 88))
                # 单元格内联样式（列宽+样式合并为单个 style，避免重复属性）
                styles: List[str] = [f"width:{w}px"]
                try:
                    if cell.fill is not None and cell.fill.fill_type == "solid":
                        bg = _xlsx_color(cell.fill)
                        if bg:
                            styles.append(f"background:{bg}")
                    if cell.font is not None:
                        if cell.font.bold:
                            styles.append("font-weight:600")
                        if cell.font.size:
                            styles.append(f"font-size:{round(cell.font.size)}px")
                        if cell.font.name:
                            styles.append(f"font-family:{cell.font.name},sans-serif")
                        fc = _xlsx_color(cell.font)
                        if fc:
                            styles.append(f"color:{fc}")
                    if cell.alignment is not None:
                        if cell.alignment.horizontal:
                            styles.append(f"text-align:{cell.alignment.horizontal}")
                        if cell.alignment.vertical:
                            styles.append(f"vertical-align:{cell.alignment.vertical}")
                except Exception:
                    pass
                attrs.append('style="' + ';'.join(styles) + '"')
                span = merge_span.get(coord)
                if span:
                    attrs.append(f'rowspan="{span[0]}"')
                    attrs.append(f'colspan="{span[1]}"')
                if cell.row in row_h:
                    attrs.append(f'height="{row_h[cell.row]}"')
                # 公式无缓存值(未重算)：计算引擎值优先，公式文本兜底
                val = cell.value
                if val is None and wb_formula is not None:
                    fv = wb_formula[ws.title].cell(
                        row=cell.row, column=cell.column).value
                    if isinstance(fv, str) and fv.startswith("="):
                        calc = _recalc.get(ws.title, {}).get(
                            str(cell.row), {}).get(get_column_letter(cell.column))
                        val = calc if calc is not None \
                            else (fv[1:] if fv.startswith("==") else fv)
                val = _fmt_value(val, cell.number_format or "")
                tag = "th" if cell.row == 1 else "td"
                attrs_str = " " + " ".join(attrs) if attrs else ""
                cells.append(f"<{tag}{attrs_str}>{html.escape(val)}</{tag}>")
            rows_html.append("<tr>" + "".join(cells) + "</tr>")
        sheet_htmls.append("<table>" + "".join(rows_html) + "</table>")
    return names, sheet_htmls


# ---- xls（xlrd） ----

def _render_xls(path: Path) -> Tuple[List[str], List[str]]:
    import xlrd
    wb = xlrd.open_workbook(str(path), formatting_info=True)
    names: List[str] = []
    sheet_htmls: List[str] = []
    for sh in wb.sheets()[:10]:
        names.append(sh.name)
        # merged cells 0-based (row_low, row_high, col_low, col_high)
        merged_covered = set()
        merge_span: Dict[Tuple[int, int], Tuple[int, int]] = {}
        for r0, r1, c0, c1 in sh.merged_cells:
            merge_span[(r0, c0)] = (r1 - r0, c1 - c0)
            for r in range(r0, r1):
                for c in range(c0, c1):
                    if (r, c) != (r0, c0):
                        merged_covered.add((r, c))
        rows_html: List[str] = []
        for r in range(sh.nrows):
            cells: List[str] = []
            for c in range(sh.ncols):
                if (r, c) in merged_covered:
                    continue
                attrs: List[str] = []
                span = merge_span.get((r, c))
                if span:
                    attrs.append(f'rowspan="{span[0]}"')
                    attrs.append(f'colspan="{span[1]}"')
                # 列宽（xlrd width 为 1/256 字符单位）
                colinfo = sh.colinfo_map.get(c)
                if colinfo and colinfo.width:
                    w = max(40.0, round((colinfo.width / 256) * 7) + 12)
                    attrs.append(f'style="width:{int(w)}px"')
                # 样式尽力提取
                style = ""
                try:
                    xf = wb.xf_list[sh.cell_xf_index(r, c)]
                    fnt = wb.font_list[xf.font_index]
                    if getattr(fnt, "bold", 0):
                        style += "font-weight:600;"
                    if getattr(fnt, "colour_index", 0) in range(1, 65):
                        style += f"color:{_xlrd_colour(wb, fnt.colour_index)};"
                    if getattr(xf, "background") and xf.background.pattern_colour_index in range(1, 65):
                        style += (f"background:"
                                  f"{_xlrd_colour(wb, xf.background.pattern_colour_index)};")
                except Exception:
                    pass
                if style:
                    attrs.append(f'style="{style}"')
                val = sh.cell_value(r, c)
                try:
                    if sh.cell_type(r, c) == xlrd.XL_CELL_DATE:
                        val = xlrd.xldate_as_datetime(val, wb.datemode)
                except Exception:
                    pass
                attrs_str = " " + " ".join(attrs) if attrs else ""
                rows_html.append(f"<td{attrs_str}>{html.escape(_fmt_value(val))}</td>")
            rows_html.append("</tr>")
        sheet_htmls.append("<table>" + "".join(rows_html) + "</table>")
    return names, sheet_htmls


def _xlrd_colour(wb, idx: int) -> str:
    """xlrd 调色板索引 → css 颜色（尽力：默认调色板前 8 色 + RGB）"""
    try:
        c = wb.colour_map.get(idx)
        if isinstance(c, (tuple, list)) and len(c) >= 4:
            r, g, b = c[0], c[1], c[2]
            return f"#{r:02x}{g:02x}{b:02x}"
    except Exception:
        pass
    basic = ["#000000", "#ffffff", "#ff0000", "#00ff00", "#0000ff",
             "#ffff00", "#00ffff", "#ff00ff"]
    return basic[idx] if 0 <= idx < len(basic) else "#000000"


# ---- csv ----

def _render_csv(path: Path) -> Tuple[List[str], List[str]]:
    from backend.services.spreadsheet.csv import _detect_encoding, _detect_delimiter
    raw = path.read_bytes()
    text = raw.decode(_detect_encoding(raw), errors="replace")
    text = text.lstrip("﻿")
    delim = _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    cells: List[str] = []
    for row in reader:
        cells.append("<tr>" + "".join(f"<td>{html.escape(c.strip())}</td>"
                                      for c in row) + "</tr>")
    return [path.stem or "CSV"], ["<table>" + "".join(cells) + "</table>"]


# ---- 入口 ----

def render_spreadsheet_html(path: Path) -> str:
    """读取表格文件 → "类 Excel" HTML（多 sheet 纯 CSS tab 切换）"""
    ext = path.suffix.lower().lstrip(".")
    if ext == "xlsx":
        names, sheet_htmls = _render_xlsx(path)
    elif ext == "xls":
        names, sheet_htmls = _render_xls(path)
    elif ext == "csv":
        names, sheet_htmls = _render_csv(path)
    else:
        raise ValueError(f"不支持的表格文件类型: .{ext}")
    if not names:
        raise RuntimeError("表格无可用工作表")
    return _HEAD_TEMPLATE.format(css=_CSS, tabs=_sheet_tabs_html(names),
                                 sheets="\n".join(
                                     _wrap_sheet(n, h, i)
                                     for i, (n, h) in enumerate(zip(names, sheet_htmls))))
